#!/usr/bin/env python3
"""Flatten tests/fixtures/SIMULATED_PRIORS_GROUND_TRUTH.json into a readable spreadsheet"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
GROUND_TRUTH_PATH = ROOT / "tests" / "fixtures" / "SIMULATED_PRIORS_GROUND_TRUTH.json"
OUTPUT_PATH = ROOT / "tests" / "fixtures" / "SIMULATED_PRIORS_GROUND_TRUTH.xlsx"

COLUMNS = [
    ("Document", 34),
    ("Type", 30),
    ("Category", 20),
    ("Prior", 40),
    ("Current", 40),
    ("Basis", 30),
    ("Expected classification", 22),
]


def _rows(ground_truth: dict):
    for document, doc_truth in ground_truth["documents"].items():
        for entry in doc_truth.get("documented_reversions", []):
            yield [
                document,
                "documented_reversion",
                entry.get("category", ""),
                entry.get("prior", ""),
                entry.get("current", ""),
                entry.get("basis", ""),
                entry.get("expected_classification", ""),
            ]
        for type_name in ("injected_unlisted_substantive_change", "injected_nonsubstantive_wording_change"):
            entry = doc_truth.get(type_name)
            if entry:
                yield [
                    document,
                    type_name,
                    entry.get("category", ""),
                    entry.get("prior", ""),
                    entry.get("current", ""),
                    entry.get("basis", ""),
                    entry.get("expected_classification", ""),
                ]


def main() -> int:
    ground_truth = json.loads(GROUND_TRUTH_PATH.read_text())

    wb = Workbook()
    ws = wb.active
    ws.title = "ground_truth"

    ws.cell(row=1, column=1, value=ground_truth.get("note", "")).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws.row_dimensions[1].height = 60

    header_row = 2
    for col_index, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=name)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width

    row = header_row + 1
    for values in _rows(ground_truth):
        for col_index, value in enumerate(values, start=1):
            ws.cell(row=row, column=col_index, value=value).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(row=header_row, column=len(COLUMNS)).column_letter}{row - 1}"

    wb.save(OUTPUT_PATH)
    print(f"OK      wrote {OUTPUT_PATH.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
