#!/usr/bin/env python3
"""Build a human-readable rendering of the output for manual QC"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent.parent
CHANGE_RESULTS_DIR = ROOT / "data" / "change_results"
OUTPUT_PATH = CHANGE_RESULTS_DIR / "change_detection_qc.xlsx"

DOCUMENT_IDS = ["mri_ct_site_of_service", "spinraza", "sleep_studies"]

CANDIDATE_COLUMNS = [
    ("change_id", 16),
    ("change_type", 11),
    ("section", 16),
    ("before_text", 45),
    ("after_text", 45),
    ("semantic_similarity", 12),
    ("revision_history_match", 12),
    ("revision_history_evidence", 40),
    ("classification", 14),
    ("changed_dimensions", 22),
    ("summary", 40),
    ("reason", 45),
    ("confidence", 10),
    ("review_status", 12),
]

_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SECTION_FONT = Font(bold=True, size=12)
_SUBSTANTIVE_FILL = PatternFill("solid", fgColor="FFC7CE")
_NON_SUBSTANTIVE_FILL = PatternFill("solid", fgColor="E2EFDA")
_UNCERTAIN_FILL = PatternFill("solid", fgColor="FFEB9C")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")

_CLASSIFICATION_FILL = {
    "substantive": _SUBSTANTIVE_FILL,
    "non_substantive": _NON_SUBSTANTIVE_FILL,
    "uncertain": _UNCERTAIN_FILL,
}


def build_sheet(ws: Worksheet, document_id: str, records: list[dict]) -> None:
    ws.cell(row=1, column=1, value=f"Detected candidates -- {document_id} ({len(records)} total)").font = _SECTION_FONT

    header_row = 2
    for col_index, (name, width) in enumerate(CANDIDATE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(col_index)].width = width

    priority = {"substantive": 0, "uncertain": 0, "non_substantive": 1}
    sorted_records = sorted(records, key=lambda r: (priority.get(r["classification"], 0), -r["confidence"]))
    class_col = next(i for i, (f, _w) in enumerate(CANDIDATE_COLUMNS, start=1) if f == "classification")

    row = header_row + 1
    for record in sorted_records:
        for col_index, (field, _width) in enumerate(CANDIDATE_COLUMNS, start=1):
            value = record.get(field)
            if isinstance(value, list):
                value = ", ".join(value)
            ws.cell(row=row, column=col_index, value=value).alignment = _WRAP_TOP
        fill = _CLASSIFICATION_FILL.get(record["classification"])
        if fill:
            ws.cell(row=row, column=class_col).fill = fill
        row += 1

    last_col_letter = get_column_letter(len(CANDIDATE_COLUMNS))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{row - 1}"
    ws.freeze_panes = f"A{header_row + 1}"


def main() -> int:
    wb = Workbook()
    wb.remove(wb.active)

    for document_id in DOCUMENT_IDS:
        records = json.loads((CHANGE_RESULTS_DIR / f"{document_id}_changes.json").read_text())
        ws = wb.create_sheet(title=document_id[:31])
        build_sheet(ws, document_id, records)

    wb.save(OUTPUT_PATH)
    print(f"OK      wrote {OUTPUT_PATH.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
