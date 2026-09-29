#!/usr/bin/env python3
"""Human-readable Excel rendering of data/review/review_queue.json"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
REVIEW_DIR = ROOT / "data" / "review"
INPUT_PATH = REVIEW_DIR / "review_queue.json"
OUTPUT_PATH = REVIEW_DIR / "review_queue.xlsx"

COLUMNS = [
    ("change_id", 16),
    ("document_id", 22),
    ("classification", 14),
    ("confidence", 10),
    ("review_status", 12),
    ("section", 16),
    ("change_type", 11),
    ("before_text", 42),
    ("after_text", 42),
    ("summary", 36),
    ("reason", 42),
    ("why_it_may_matter", 36),
    ("billing_setting", 16),
    ("service_area", 14),
    ("age_min", 8),
    ("age_max", 8),
    ("codes", 14),
    ("states", 12),
    ("plan_scope", 14),
    ("additional_data_needed", 26),
    ("revision_history_match", 12),
    ("revision_history_evidence", 36),
    ("semantic_similarity", 12),
    ("effective_date", 12),
    ("source_url", 30),
    ("prior_policy_number", 16),
    ("current_policy_number", 16),
    ("prior_is_simulated", 12),
    ("prior_page", 10),
    ("current_page", 10),
    ("changed_dimensions", 24),
]

_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SECTION_FONT = Font(bold=True, size=12)
_SUBSTANTIVE_FILL = PatternFill("solid", fgColor="FFC7CE")
_NON_SUBSTANTIVE_FILL = PatternFill("solid", fgColor="E2EFDA")
_UNCERTAIN_FILL = PatternFill("solid", fgColor="FFEB9C")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")

_CLASSIFICATION_FILL = {
    "substantive": _SUBSTANTIVE_FILL,
    "non_substantive": _NON_SUBSTANTIVE_FILL,
    "uncertain": _UNCERTAIN_FILL,
}


def main() -> int:
    records = json.loads(INPUT_PATH.read_text())

    wb = Workbook()
    ws = wb.active
    ws.title = "review_queue"

    ws.cell(row=1, column=1, value=f"data/review/review_queue.json -- {len(records)} total candidates").font = (
        _SECTION_FONT
    )

    header_row = 2
    for col_index, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(col_index)].width = width

    class_col = next(i for i, (f, _w) in enumerate(COLUMNS, start=1) if f == "classification")

    row = header_row + 1
    for record in records:
        for col_index, (field, _width) in enumerate(COLUMNS, start=1):
            value = record.get(field)
            if isinstance(value, list):
                value = ", ".join(value)
            ws.cell(row=row, column=col_index, value=value).alignment = _WRAP_TOP
        fill = _CLASSIFICATION_FILL.get(record.get("classification"))
        if fill:
            ws.cell(row=row, column=class_col).fill = fill
        row += 1

    last_col_letter = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{row - 1}"
    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(OUTPUT_PATH)
    print(f"OK      wrote {OUTPUT_PATH.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
